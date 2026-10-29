import unicodedata
from io import BytesIO

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from database import db, sincronizar_actividades_ordenes_abiertas
from auth import require_admin
from routers.configuracion import sincronizar_secuencia, texto_obligatorio, entero_obligatorio

router = APIRouter()

JORNADA_MIN = 535  # misma jornada que usan las hojas de tiempos del Excel


def no_autorizado():
    return JSONResponse({"detail": "No autorizado"}, status_code=401)


def minutos_validos(valor, obligatorio=True):
    if valor is None or str(valor).strip() == "":
        if obligatorio:
            raise HTTPException(status_code=400, detail="Los minutos por unidad son obligatorios")
        return None
    try:
        minutos = float(str(valor).strip().replace(",", "."))
    except Exception:
        raise HTTPException(status_code=400, detail="Minutos por unidad invalidos")
    if minutos <= 0:
        raise HTTPException(status_code=400, detail="Los minutos por unidad deben ser mayores a 0")
    return minutos


def costo_valido(valor):
    if valor is None or str(valor).strip() == "":
        return 0
    try:
        costo = float(str(valor).strip().replace(",", "."))
    except Exception:
        raise HTTPException(status_code=400, detail="Costo invalido")
    if costo < 0:
        raise HTTPException(status_code=400, detail="El costo no puede ser negativo")
    return costo


def operario_valido(cursor, valor):
    if valor in (None, "", 0, "0"):
        return None
    try:
        operario_id = int(valor)
    except Exception:
        raise HTTPException(status_code=400, detail="Operario invalido")
    cursor.execute("SELECT id FROM operarios WHERE id = %s", (operario_id,))
    if not cursor.fetchone():
        raise HTTPException(status_code=400, detail="El operario no existe")
    return operario_id


def texto_corto(valor, maximo=120):
    texto = str(valor or "").strip()
    return texto[:maximo] if texto else None


def guardar_estandar(cursor, actividad_id: int, minutos: float, costo_mo_unidad=None):
    unidades_por_hora = 60.0 / minutos
    cursor.execute("SELECT id FROM estandares_actividad WHERE actividad_id = %s ORDER BY id LIMIT 1", (actividad_id,))
    row = cursor.fetchone()
    if row:
        if costo_mo_unidad is None:
            cursor.execute(
                "UPDATE estandares_actividad SET unidades_por_hora = %s WHERE id = %s",
                (unidades_por_hora, row[0]),
            )
        else:
            cursor.execute(
                "UPDATE estandares_actividad SET unidades_por_hora = %s, costo_mo_unidad = %s WHERE id = %s",
                (unidades_por_hora, costo_mo_unidad, row[0]),
            )
        return row[0]
    sincronizar_secuencia(cursor, "estandares_actividad")
    cursor.execute(
        """
        INSERT INTO estandares_actividad (actividad_id, unidades_por_hora, costo_mo_unidad, costo_mo_hora)
        VALUES (%s, %s, %s, 0) RETURNING id
        """,
        (actividad_id, unidades_por_hora, costo_mo_unidad or 0),
    )
    return cursor.fetchone()[0]


def leer_tiempos(cursor, maquina_id: int):
    cursor.execute("SELECT id, nombre, COALESCE(activo, TRUE) FROM maquinas WHERE id = %s", (maquina_id,))
    maquina = cursor.fetchone()
    if not maquina:
        raise HTTPException(status_code=404, detail="Producto no encontrado")

    cursor.execute(
        """
        SELECT id, nombre, COALESCE(activo, TRUE)
        FROM procesos WHERE maquina_id = %s ORDER BY id
        """,
        (maquina_id,),
    )
    procesos = [{"id": r[0], "nombre": r[1], "activo": r[2], "actividades": []} for r in cursor.fetchall()]
    por_id = {p["id"]: p for p in procesos}

    cursor.execute(
        """
        SELECT a.id, a.proceso_id, a.nombre, COALESCE(a.activo, TRUE),
               a.herramienta, a.operario_responsable_id,
               e.unidades_por_hora, e.costo_mo_unidad,
               r.mediana, COALESCE(r.registros, 0)
        FROM actividades a
        JOIN procesos p ON p.id = a.proceso_id
        LEFT JOIN LATERAL (
            SELECT unidades_por_hora, costo_mo_unidad
            FROM estandares_actividad WHERE actividad_id = a.id ORDER BY id LIMIT 1
        ) e ON TRUE
        LEFT JOIN LATERAL (
            SELECT percentile_cont(0.5) WITHIN GROUP (
                       ORDER BY (EXTRACT(EPOCH FROM (rp.fin::timestamp - rp.inicio::timestamp)) / 60.0) / rp.cantidad
                   ) AS mediana,
                   COUNT(*) AS registros
            FROM registros_produccion rp
            WHERE rp.actividad_id = a.id AND rp.cantidad > 0 AND rp.fin IS NOT NULL
              AND rp.fin::timestamp > rp.inicio::timestamp
        ) r ON TRUE
        WHERE p.maquina_id = %s
        ORDER BY a.proceso_id, a.id
        """,
        (maquina_id,),
    )
    for r in cursor.fetchall():
        uph = float(r[6]) if r[6] else None
        por_id[r[1]]["actividades"].append({
            "id": r[0],
            "nombre": r[2],
            "activo": r[3],
            "herramienta": r[4] or "",
            "operario_id": r[5],
            "min_unidad": round(60.0 / uph, 2) if uph else None,
            "costo_mo_unidad": round(float(r[7]), 2) if r[7] is not None else 0,
            "real_mediana": round(float(r[8]), 2) if r[8] is not None else None,
            "real_registros": int(r[9]),
        })
    return {"id": maquina[0], "nombre": maquina[1], "activo": maquina[2]}, procesos


@router.get("/configuracion/tiempos", response_class=HTMLResponse)
def vista_tiempos(request: Request):
    if "username" not in request.session:
        return RedirectResponse("/admin?next=configuracion/tiempos", 303)
    if request.session.get("role") != "admin":
        return RedirectResponse("/", 303)
    return request.app.state.templates.TemplateResponse(
        request=request, name="config_tiempos.html", context={"request": request}
    )


@router.get("/config/tiempos")
def datos_tiempos(request: Request, maquina_id: int = 0):
    if not require_admin(request):
        return no_autorizado()
    conn = db()
    c = conn.cursor()
    try:
        c.execute("SELECT id, nombre, COALESCE(activo, TRUE) FROM maquinas ORDER BY nombre, id")
        maquinas = [{"id": r[0], "nombre": r[1], "activo": r[2]} for r in c.fetchall()]
        c.execute("SELECT id, nombre FROM operarios WHERE COALESCE(activo, TRUE) = TRUE ORDER BY nombre")
        operarios = [{"id": r[0], "nombre": r[1]} for r in c.fetchall()]
        maquina, procesos = (None, [])
        if maquina_id:
            maquina, procesos = leer_tiempos(c, maquina_id)
    finally:
        conn.close()
    return {"maquinas": maquinas, "operarios": operarios, "maquina": maquina,
            "procesos": procesos, "jornada_min": JORNADA_MIN}


@router.post("/config/tiempos/proceso")
def crear_proceso(data: dict, request: Request):
    if not require_admin(request):
        return no_autorizado()
    maquina_id = entero_obligatorio(data, "maquina_id", "Producto")
    nombre = texto_obligatorio(data, "nombre", "Nombre del proceso")
    conn = db()
    c = conn.cursor()
    try:
        c.execute("SELECT id FROM maquinas WHERE id = %s", (maquina_id,))
        if not c.fetchone():
            raise HTTPException(status_code=400, detail="El producto no existe")
        c.execute(
            "SELECT id FROM procesos WHERE maquina_id = %s AND lower(trim(nombre)) = lower(trim(%s))",
            (maquina_id, nombre),
        )
        if c.fetchone():
            raise HTTPException(status_code=400, detail="Ese proceso ya existe en este producto")
        sincronizar_secuencia(c, "procesos")
        c.execute("INSERT INTO procesos (maquina_id, nombre) VALUES (%s, %s) RETURNING id", (maquina_id, nombre))
        proceso_id = c.fetchone()[0]
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return {"ok": True, "id": proceso_id}


@router.put("/config/tiempos/proceso/{proceso_id}")
def editar_proceso(proceso_id: int, data: dict, request: Request):
    if not require_admin(request):
        return no_autorizado()
    conn = db()
    c = conn.cursor()
    try:
        c.execute("SELECT id FROM procesos WHERE id = %s", (proceso_id,))
        if not c.fetchone():
            raise HTTPException(status_code=404, detail="Proceso no encontrado")
        if "nombre" in data:
            c.execute("UPDATE procesos SET nombre = %s WHERE id = %s",
                      (texto_obligatorio(data, "nombre", "Nombre del proceso"), proceso_id))
        if "activo" in data:
            c.execute("UPDATE procesos SET activo = %s WHERE id = %s", (bool(data["activo"]), proceso_id))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return {"ok": True}


@router.post("/config/tiempos/actividad")
def crear_actividad(data: dict, request: Request):
    if not require_admin(request):
        return no_autorizado()
    proceso_id = entero_obligatorio(data, "proceso_id", "Proceso")
    nombre = texto_obligatorio(data, "nombre", "Nombre de la actividad")
    minutos = minutos_validos(data.get("min_unidad"))
    costo = costo_valido(data.get("costo_mo_unidad"))
    conn = db()
    c = conn.cursor()
    try:
        c.execute("SELECT id FROM procesos WHERE id = %s", (proceso_id,))
        if not c.fetchone():
            raise HTTPException(status_code=400, detail="El proceso no existe")
        c.execute(
            "SELECT id FROM actividades WHERE proceso_id = %s AND lower(trim(nombre)) = lower(trim(%s))",
            (proceso_id, nombre),
        )
        if c.fetchone():
            raise HTTPException(status_code=400, detail="Esa actividad ya existe en este proceso")
        operario_id = operario_valido(c, data.get("operario_id"))
        sincronizar_secuencia(c, "actividades")
        c.execute(
            """
            INSERT INTO actividades (proceso_id, nombre, herramienta, operario_responsable_id)
            VALUES (%s, %s, %s, %s) RETURNING id
            """,
            (proceso_id, nombre, texto_corto(data.get("herramienta")), operario_id),
        )
        actividad_id = c.fetchone()[0]
        guardar_estandar(c, actividad_id, minutos, costo)
        ordenes = sincronizar_actividades_ordenes_abiertas(c)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return {"ok": True, "id": actividad_id, "ordenes_actualizadas": ordenes}


@router.put("/config/tiempos/actividad/{actividad_id}")
def editar_actividad(actividad_id: int, data: dict, request: Request):
    if not require_admin(request):
        return no_autorizado()
    conn = db()
    c = conn.cursor()
    try:
        c.execute("SELECT id FROM actividades WHERE id = %s", (actividad_id,))
        if not c.fetchone():
            raise HTTPException(status_code=404, detail="Actividad no encontrada")
        if "nombre" in data:
            c.execute("UPDATE actividades SET nombre = %s WHERE id = %s",
                      (texto_obligatorio(data, "nombre", "Nombre de la actividad"), actividad_id))
        if "herramienta" in data:
            c.execute("UPDATE actividades SET herramienta = %s WHERE id = %s",
                      (texto_corto(data.get("herramienta")), actividad_id))
        if "operario_id" in data:
            c.execute("UPDATE actividades SET operario_responsable_id = %s WHERE id = %s",
                      (operario_valido(c, data.get("operario_id")), actividad_id))
        if "activo" in data:
            c.execute("UPDATE actividades SET activo = %s WHERE id = %s", (bool(data["activo"]), actividad_id))
        if "min_unidad" in data or "costo_mo_unidad" in data:
            minutos = minutos_validos(data.get("min_unidad"), obligatorio=False)
            costo = costo_valido(data["costo_mo_unidad"]) if "costo_mo_unidad" in data else None
            if minutos is None:
                c.execute("SELECT unidades_por_hora FROM estandares_actividad WHERE actividad_id = %s ORDER BY id LIMIT 1",
                          (actividad_id,))
                row = c.fetchone()
                if not row or not row[0]:
                    raise HTTPException(status_code=400, detail="Primero ingrese los minutos por unidad")
                minutos = 60.0 / float(row[0])
            guardar_estandar(c, actividad_id, minutos, costo)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return {"ok": True}


@router.get("/config/tiempos/exportar")
def exportar_tiempos(request: Request, maquina_id: int):
    if not require_admin(request):
        return no_autorizado()
    conn = db()
    c = conn.cursor()
    try:
        maquina, procesos = leer_tiempos(c, maquina_id)
        c.execute("SELECT id, nombre FROM operarios")
        operarios = {r[0]: r[1] for r in c.fetchall()}
    finally:
        conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "tiempos"
    negrita = Font(name="Arial", bold=True)
    normal = Font(name="Arial")
    verde = PatternFill("solid", fgColor="2E7D32")
    ws["B1"] = f"Tiempos {maquina['nombre']}"
    ws["B1"].font = Font(name="Arial", bold=True, size=13)
    encabezados = ["ITEM", "PROCESO", "SUBPROCESO", "min/unidad", "Maquina", "Operario", "Real mediana (min)", "Registros"]
    for col, texto in enumerate(encabezados, start=2):
        celda = ws.cell(row=3, column=col, value=texto)
        celda.font = Font(name="Arial", bold=True, color="FFFFFF")
        celda.fill = verde
        celda.alignment = Alignment(horizontal="center", wrap_text=True)

    fila = 4
    item = 1
    for proceso in procesos:
        if not proceso["activo"]:
            continue
        primera = True
        for act in proceso["actividades"]:
            if not act["activo"]:
                continue
            valores = [item, proceso["nombre"] if primera else None, act["nombre"], act["min_unidad"],
                       act["herramienta"] or None, operarios.get(act["operario_id"]),
                       act["real_mediana"], act["real_registros"] or None]
            for col, valor in enumerate(valores, start=2):
                ws.cell(row=fila, column=col, value=valor).font = normal
            ws.cell(row=fila, column=5).number_format = "0.00"
            ws.cell(row=fila, column=8).number_format = "0.00"
            primera = False
            item += 1
            fila += 1

    ultima = fila - 1
    total = fila
    ws.cell(row=total, column=4, value="Tiempo de ciclo minutos").font = negrita
    ws.cell(row=total, column=5, value=f"=SUM(E4:E{ultima})" if ultima >= 4 else 0).font = negrita
    ws.cell(row=total + 1, column=4, value="Operarios").font = negrita
    ws.cell(row=total + 1, column=5, value=1).font = Font(name="Arial", color="0000FF")
    ws.cell(row=total + 2, column=4, value="Jornada laboral min").font = negrita
    ws.cell(row=total + 2, column=5, value=JORNADA_MIN).font = Font(name="Arial", color="0000FF")
    ws.cell(row=total + 3, column=4, value="Tiempo disponible x turno").font = negrita
    ws.cell(row=total + 3, column=5, value=f"=E{total + 1}*E{total + 2}").font = normal
    ws.cell(row=total + 4, column=4, value="Meta unidades x dia").font = negrita
    ws.cell(row=total + 4, column=5, value=f"=IF(E{total}>0,E{total + 3}/E{total},0)").font = normal
    for r in range(total, total + 5):
        ws.cell(row=r, column=5).number_format = "0.00"
    ws.cell(row=total + 6, column=2,
            value="Exportado desde la web de produccion. min/unidad = 60 / unidades por hora del estandar. "
                  "Operarios y jornada (en azul) se pueden ajustar.").font = Font(name="Arial", italic=True, size=9)

    for col, ancho in zip("BCDEFGHI", [7, 26, 44, 12, 30, 24, 14, 11]):
        ws.column_dimensions[col].width = ancho

    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    ascii_nombre = unicodedata.normalize("NFKD", maquina["nombre"]).encode("ascii", "ignore").decode()
    nombre = "".join(ch if ch.isalnum() or ch in " -_." else "_" for ch in ascii_nombre).strip()
    return StreamingResponse(
        salida,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="tiempos {nombre}.xlsx"'},
    )
